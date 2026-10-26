#!/usr/bin/env python3
"""Monta a planilha de acompanhamento de anulação a partir do JSON extraído.

Gera um XLSX com uma aba "Empenhos" — dados do relatório à esquerda, colunas de
preenchimento à direita, com listas de escolha — e uma aba "Resumo" com os
totais por prioridade, execução, idade e fonte.

Regras aplicadas (as mesmas da planilha compartilhada):

  saldo a anular  coluna "À Liquidar" do relatório, isto é
                  empenhado + reforçado - anulado - liquidado. Só essa parcela
                  pode ser anulada; o que já foi liquidado é direito do credor.
  execução        "Sem execução" se nada foi liquidado; senão "Sobra de saldo".
  prioridade      Alta  = sem execução há mais de 90 dias
                  Média = demais empenhos com mais de 60 dias
                  Baixa = até 60 dias (ainda no fluxo normal de execução)

Uso:
    python3 gerar_planilha_xlsx.py empenhos.json [-o planilha.xlsx]
                                   [--referencia DD/MM/AAAA]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DECISOES = ["Pendente", "Anular integral", "Anular parcial", "Manter", "Aguardar informação"]
SITUACOES = ["Em triagem", "Com parecer do setor", "Aprovado para anulação",
             "Anulado no sistema", "Arquivado"]
SETORES = ["Execução Orçamentária e Financeira", "Compras e Licitações", "Almoxarifado",
           "Gestor do contrato", "Controle Interno", "Gabinete da Secretaria"]

# Natureza da despesa pelos quatro primeiros grupos do elemento, quando o
# relatório inteiro é de uma só natureza (é o caso dos relatórios por tipo).
NATUREZAS = {
    "3.3.90.30": "Material de consumo",
    "3.3.90.32": "Material de distribuição gratuita",
    "3.3.90.36": "Serviços de terceiros - pessoa física",
    "3.3.90.39": "Serviços de terceiros - pessoa jurídica",
    "3.3.90.40": "Serviços de tecnologia da informação",
    "4.4.90.51": "Obras e instalações",
    "4.4.90.52": "Equipamentos e material permanente",
}
# Modalidades de empenho (Lei 4.320/64, art. 60)
TIPOS = {"OR": "Ordinário", "ES": "Estimativo", "GL": "Global"}


def natureza_de(empenhos: list[dict]) -> str:
    """Nome da natureza quando todos os empenhos partilham o mesmo elemento."""
    elementos = {".".join(r["categoria"].split(".")[:4]) for r in empenhos}
    if len(elementos) != 1:
        return ""
    elemento = elementos.pop()
    nome = NATUREZAS.get(elemento)
    return f"{nome} ({elemento})" if nome else elemento

TINTA = "0D5C66"
PREENCHER = "FFF7E0"          # amarelo claro: colunas que a equipe preenche

CABECALHO = [
    ("Empenho", 9), ("Tipo", 6), ("Data", 11), ("Idade (dias)", 12), ("Fornecedor", 42),
    ("Ficha", 7), ("Vínculo", 10), ("Fonte", 11), ("Unid. orç.", 11), ("Funcional", 22),
    ("Elemento", 13), ("Empenhado", 14), ("Anulado", 11), ("Liquidado", 14), ("Pago", 14),
    ("Saldo a anular", 15), ("Execução", 15), ("Prioridade sugerida", 18),
    ("Decisão", 18), ("Valor a anular", 14), ("Setor responsável", 30),
    ("Justificativa / parecer", 46), ("Processo / documento", 22), ("Situação", 22),
    ("Responsável", 20), ("Data da análise", 16),
]
COL_DECISAO, COL_VALOR, COL_SETOR, COL_SITUACAO = 19, 20, 21, 24
COL_SALDO = 16


def dec(s: str) -> Decimal:
    return Decimal(s.replace(".", "").replace(",", "."))


def faixa_idade(d: int) -> str:
    if d > 180:
        return ">180 dias"
    if d > 90:
        return "91-180 dias"
    if d > 60:
        return "61-90 dias"
    if d > 30:
        return "31-60 dias"
    return "Até 30 dias"


def prepara(empenhos: list[dict], referencia: dt.date) -> list[tuple]:
    saida = []
    for r in empenhos:
        data = dt.datetime.strptime(r["data"], "%d/%m/%Y").date()
        idade = (referencia - data).days
        liquidado = dec(r["liquidado"])
        execucao = "Sem execução" if liquidado == 0 else "Sobra de saldo"
        if liquidado == 0 and idade > 90:
            prioridade = "Alta"
        elif idade > 60:
            prioridade = "Média"
        else:
            prioridade = "Baixa"
        saida.append((r, data, idade, execucao, prioridade))
    saida.sort(key=lambda x: -x[2])          # mais antigos primeiro
    return saida


def aba_empenhos(wb: Workbook, linhas: list[tuple]) -> None:
    ws = wb.active
    ws.title = "Empenhos"
    f_tit = Font(bold=True, color="FFFFFF")
    f_neg = Font(bold=True)
    fill_tit = PatternFill("solid", fgColor=TINTA)
    fill_pre = PatternFill("solid", fgColor=PREENCHER)
    borda = Border(bottom=Side("thin", color="D2D9DD"))

    ws.append([c for c, _ in CABECALHO])
    for i, (_, larg) in enumerate(CABECALHO, 1):
        cel = ws.cell(row=1, column=i)
        cel.font, cel.fill = f_tit, fill_tit
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 30

    for r, data, idade, execucao, prioridade in linhas:
        ws.append([
            r["emp"], r["tipo"], data, idade, r["fornecedor"], r["ficha"], r["vinculo"],
            r["fonte"], r["unid_orc"], r["funcional"], r["categoria"],
            float(dec(r["empenhado"])), float(dec(r["anulado"])), float(dec(r["liquidado"])),
            float(dec(r["pago"])), float(dec(r["a_liquidar"])), execucao, prioridade,
            "Pendente", None, None, None, None, "Em triagem", None, None,
        ])

    ultima = ws.max_row
    for linha in ws.iter_rows(min_row=2, max_row=ultima, max_col=len(CABECALHO)):
        for cel in linha:
            cel.border = borda
            if cel.column in (12, 13, 14, 15, COL_SALDO, COL_VALOR):
                cel.number_format = "#,##0.00"
            elif cel.column in (3, 26):
                cel.number_format = "DD/MM/YYYY"
            elif cel.column in (5, 22):
                cel.alignment = Alignment(wrap_text=True, vertical="top")
            if cel.column >= COL_DECISAO:
                cel.fill = fill_pre
        ws.cell(row=linha[0].row, column=COL_SALDO).font = f_neg

    def lista(coluna: int, opcoes: list[str], titulo: str) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        dv.errorTitle, dv.error = titulo, "Escolha uma das opções da lista."
        ws.add_data_validation(dv)
        letra = get_column_letter(coluna)
        dv.add(f"{letra}2:{letra}{ultima}")

    lista(COL_DECISAO, DECISOES, "Decisão")
    lista(COL_SETOR, SETORES, "Setor responsável")
    lista(COL_SITUACAO, SITUACOES, "Situação")

    # o valor a anular nunca passa do saldo a anular da própria linha
    letra_v, letra_s = get_column_letter(COL_VALOR), get_column_letter(COL_SALDO)
    dv = DataValidation(type="decimal", operator="between", formula1="0",
                        formula2=f"${letra_s}2", allow_blank=True)
    dv.errorTitle = "Valor a anular"
    dv.error = "O valor a anular não pode passar do saldo a anular desta linha."
    ws.add_data_validation(dv)
    dv.add(f"{letra_v}2:{letra_v}{ultima}")

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{ultima}"


def aba_resumo(wb: Workbook, linhas: list[tuple], cab: dict, referencia: dt.date,
               natureza: str = "") -> None:
    rs = wb.create_sheet("Resumo")
    f_tit = Font(bold=True, color="FFFFFF")
    f_neg = Font(bold=True)
    fill_tit = PatternFill("solid", fgColor=TINTA)
    for col, larg in (("A", 34), ("B", 11), ("C", 18)):
        rs.column_dimensions[col].width = larg

    rs.cell(row=1, column=1, value=cab.get("orgao", "")).font = f_neg
    rs.cell(row=2, column=1,
            value=f'Listagem de empenhos, situação em {cab.get("situacao", "")}, '
                  f'emitida em {cab.get("emissao", "")}.')
    rs.cell(row=3, column=1,
            value=f'Idade contada até {referencia.strftime("%d/%m/%Y")}. '
                  f'Saldo a anular = coluna "À Liquidar".')
    linha = 4
    if natureza:
        rs.cell(row=linha, column=1, value=f"Natureza da despesa: {natureza}")
        linha += 1
    linha += 1

    def bloco(titulo: str, itens: list[tuple[str, int, Decimal]]) -> None:
        nonlocal linha
        for col, texto in ((1, titulo), (2, "Empenhos"), (3, "Saldo a anular")):
            cel = rs.cell(row=linha, column=col, value=texto)
            cel.font, cel.fill = f_tit, fill_tit
        for i, (rot, qtd, val) in enumerate(itens, 1):
            rs.cell(row=linha + i, column=1, value=rot)
            rs.cell(row=linha + i, column=2, value=qtd)
            cel = rs.cell(row=linha + i, column=3, value=float(val))
            cel.number_format = "#,##0.00"
        linha += len(itens) + 2

    def agrupa(chave):
        g: dict[str, list] = {}
        for r, data, idade, execucao, prioridade in linhas:
            k = chave(r, idade, execucao, prioridade)
            a = g.setdefault(k, [0, Decimal(0)])
            a[0] += 1
            a[1] += dec(r["a_liquidar"])
        return g

    g = agrupa(lambda r, i, e, p: p)
    bloco("Por prioridade sugerida", [(k, *g[k]) for k in ("Alta", "Média", "Baixa") if k in g])
    g = agrupa(lambda r, i, e, p: e)
    bloco("Por execução", [(k, *g[k]) for k in sorted(g)])
    g = agrupa(lambda r, i, e, p: faixa_idade(i))
    ordem = [">180 dias", "91-180 dias", "61-90 dias", "31-60 dias", "Até 30 dias"]
    bloco("Por idade do empenho", [(k, *g[k]) for k in ordem if k in g])
    g = agrupa(lambda r, i, e, p: TIPOS.get(r["tipo"], r["tipo"]))
    if len(g) > 1:
        bloco("Por modalidade de empenho", [(k, *g[k]) for k in sorted(g, key=lambda x: -g[x][1])])
    g = agrupa(lambda r, i, e, p: r["fonte"])
    bloco("Por fonte de recurso", [(k, *g[k]) for k in sorted(g, key=lambda x: -g[x][1])])

    total = sum(dec(r["a_liquidar"]) for r, *_ in linhas)
    rs.cell(row=linha, column=1, value="TOTAL").font = f_neg
    rs.cell(row=linha, column=2, value=len(linhas)).font = f_neg
    cel = rs.cell(row=linha, column=3, value=float(total))
    cel.font, cel.number_format = f_neg, "#,##0.00"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json", type=Path, help="saída de extrair_empenhos.py")
    ap.add_argument("-o", "--saida", type=Path, default=Path("planilha_anulacao.xlsx"))
    ap.add_argument("--referencia", default=None,
                    help="data base da idade (DD/MM/AAAA); por omissão, a emissão do relatório")
    ap.add_argument("--natureza", default=None,
                    help="rótulo da natureza da despesa; por omissão, deduzido do elemento")
    args = ap.parse_args()

    if not args.json.is_file():
        print(f"arquivo não encontrado: {args.json}", file=sys.stderr)
        return 2
    dados = json.loads(args.json.read_text(encoding="utf-8"))
    cab = dados.get("cabecalho", {})

    if args.referencia:
        referencia = dt.datetime.strptime(args.referencia, "%d/%m/%Y").date()
    elif cab.get("emissao"):
        referencia = dt.datetime.strptime(cab["emissao"].split()[0], "%d/%m/%Y").date()
    else:
        referencia = dt.date.today()

    natureza = args.natureza or natureza_de(dados["empenhos"])
    linhas = prepara(dados["empenhos"], referencia)
    wb = Workbook()
    aba_empenhos(wb, linhas)
    aba_resumo(wb, linhas, cab, referencia, natureza)
    wb.save(args.saida)

    total = sum(dec(r["a_liquidar"]) for r, *_ in linhas)
    print(f'{len(linhas)} empenhos · saldo a anular '
          f'{f"{total:,.2f}".replace(",", chr(0)).replace(".", ",").replace(chr(0), ".")}')
    print(f"gravado em {args.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
